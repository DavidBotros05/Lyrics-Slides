from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os
from pathlib import Path
import sys
import random

try:
    from PIL import Image
    _PIL_AVAILABLE = True
except ImportError:
    _PIL_AVAILABLE = False


CUR_DIR = str(Path(__file__).resolve().parent)

BG_IMAGE = CUR_DIR + '/background_images/test.png'  # <-- change to your image path
OUTPUT_PPTX = CUR_DIR + "/"

# ===== Text color you can tweak (RGB) =====
# Examples:
#   White: (255, 255, 255)
#   Black: (0, 0, 0)
#   Yellow: (255, 255, 0)
#   Light gray: (220, 220, 220)
TEXT_RGB = (255, 255, 255)

def _relative_luminance(r, g, b):
    """WCAG relative luminance (0..1) for an sRGB pixel."""
    def lin(c):
        c = c / 255.0
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def _contrast_ratio(l1, l2):
    """WCAG contrast ratio for two relative luminance values."""
    lighter = max(l1, l2)
    darker = min(l1, l2)
    return (lighter + 0.05) / (darker + 0.05)


def pick_text_color(image_path):
    """Pick black or white text for best contrast against the background.

    Averages the whole image and returns whichever of black/white has the
    higher WCAG contrast ratio against that average brightness.
    """
    if not _PIL_AVAILABLE:
        return None
    try:
        with Image.open(image_path) as im:
            # Transparent areas show the white slide behind the picture, so
            # composite onto white first - otherwise PNG transparency would
            # be counted as black and flip the choice the wrong way.
            if im.mode in ('RGBA', 'LA', 'PA') or (
                    im.mode == 'P' and 'transparency' in im.info):
                im = im.convert('RGBA')
                white = Image.new('RGBA', im.size, (255, 255, 255, 255))
                im = Image.alpha_composite(white, im)
            im = im.convert('RGB')
            im = im.resize((64, 64))
            pixels = list(im.getdata())
    except Exception:
        return None

    if not pixels:
        return None

    avg_l = sum(_relative_luminance(*p) for p in pixels) / len(pixels)
    contrast_white = _contrast_ratio(1.0, avg_l)
    contrast_black = _contrast_ratio(0.0, avg_l)
    return (255, 255, 255) if contrast_white >= contrast_black else (0, 0, 0)


# Per-image placement, as fractions of the slide size:
#   x / y = offset of the picture's top-left corner (0 = top-left of slide,
#           negative or >1 puts it partly outside the slide -> cropped)
#   w / h = picture size (1.0 = exactly the slide's width/height)
# The identity transform stretches the image to fill the slide exactly.
# Saved per-image layouts (from the UI's "Adjust" editor) live in
# lyrics_slides_settings.json under "background_transforms" and are passed in
# via the `transform` argument; these defaults cover images without one.
IDENTITY_TRANSFORM = {'x': 0.0, 'y': 0.0, 'w': 1.0, 'h': 1.0}
DEFAULT_TRANSFORMS = {
    # Former hardcoded placements, expressed as slide fractions.
    'image1.jpeg': {'x': 0.0, 'y': 0.0, 'w': 1.0, 'h': 20.2 / 19.05},
    'Unmute.png': {'x': 0.0, 'y': -12.64 / 19.05, 'w': 1.0, 'h': 38.1 / 19.05},
}


def Background_image(path, prs, source = 'random', color = (255,255,255), auto_color = True,
                     transform = None):
    images = []

    if source == 'random':
        for image in os.listdir(path):
            if image[0] == ',':
                continue
            if image[-4:].lower() == '.png' or image[-5:].lower() == '.jpeg' or image[-4:].lower() == '.jpg':
                images.append(image)


        number = random.randint(0,len(images)-1)
        image = images[number]


    else:
        image = source

    # Saved layout from the UI wins; otherwise a known default; otherwise
    # stretch-to-fill.
    t = dict(IDENTITY_TRANSFORM)
    t.update(DEFAULT_TRANSFORMS.get(image, {}))
    if isinstance(transform, dict):
        for key in ('x', 'y', 'w', 'h'):
            if key in transform:
                try:
                    t[key] = float(transform[key])
                except (TypeError, ValueError):
                    pass

    info_image = {'image': image, 'color': color,
                  'width': int(round(prs.slide_width * t['w'])),
                  'height': int(round(prs.slide_height * t['h'])),
                  'horizontal offset': int(round(prs.slide_width * t['x'])),
                  'vertical offset': int(round(prs.slide_height * t['y']))}

    # Auto mode: pick black/white from the image's brightness. The only
    # override is the Black/White choice in the UI (auto_color=False), in
    # which case the caller's color is used as-is.
    if auto_color:
        chosen = pick_text_color(os.path.join(path, image))
        if chosen is not None:
            info_image['color'] = chosen

    return info_image

    
       




def send_shape_to_back(slide, shape) -> None:
    """Move `shape` behind all other shapes on the slide."""
    spTree = slide.shapes._spTree  # lxml element
    el = shape._element
    spTree.remove(el)
    # Index 2 is commonly right after p:nvGrpSpPr and p:grpSpPr
    spTree.insert(2, el)


def remove_all_pictures(slide) -> int:
    """Remove all picture shapes from the slide.

    This is intentionally aggressive: if any image is already present on the slide,
    we delete it so the new background image is the only picture.

    Returns the number of removed shapes.
    """
    removed = 0

    # Iterate over a copy because we'll remove elements from the underlying XML tree.
    for shape in list(slide.shapes):
        try:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue

            slide.shapes._spTree.remove(shape._element)
            removed += 1
        except Exception:
            continue

    return removed


def _contrast_color(rgb: tuple[int, int, int]) -> tuple[int, int, int]:
    """Black for light text, white for dark text (used for outline/shadow)."""
    return (0, 0, 0) if _relative_luminance(*rgb) >= 0.5 else (255, 255, 255)


def _apply_outline_and_shadow_to_run(run, effect_rgb: tuple[int, int, int]) -> None:
    """Add a thin letter contour (a:ln) and a soft outer shadow to one run.

    python-pptx has no API for text outlines or per-run shadows, so the
    DrawingML elements are inserted directly into the run properties XML.
    This keeps lyrics readable on any background.
    """
    from lxml import etree

    a_ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    qn = lambda tag: '{%s}%s' % (a_ns, tag)

    rPr = run._r.get_or_add_rPr()

    # Remove any outline/effects we may have added on a previous pass.
    for tag in ('ln', 'effectLst'):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)

    hex_color = '%02X%02X%02X' % effect_rgb

    # 1) Letter contour: ~1pt outline in the contrast color (w is in EMU/12700ths of a point).
    ln = etree.fromstring(
        f'<a:ln xmlns:a="{a_ns}" w="12700" cap="rnd">'
        f'<a:solidFill><a:srgbClr val="{hex_color}"/></a:solidFill>'
        f'<a:round/></a:ln>')
    # Schema order: a:ln must be the first child element of a:rPr.
    rPr.insert(0, ln)

    # 2) Soft outer shadow / halo in the same contrast color.
    effect = etree.fromstring(
        f'<a:effectLst xmlns:a="{a_ns}">'
        f'<a:outerShdw blurRad="63500" dist="38100" dir="2700000" rotWithShape="0">'
        f'<a:srgbClr val="{hex_color}"><a:alpha val="70000"/></a:srgbClr>'
        f'</a:outerShdw></a:effectLst>')
    # Schema order: a:effectLst goes right after the fill element(s).
    fill_tags = {qn(t) for t in
                 ('ln', 'noFill', 'solidFill', 'gradFill', 'blipFill',
                  'pattFill', 'grpFill')}
    insert_at = 0
    for i, child in enumerate(rPr):
        if child.tag in fill_tags:
            insert_at = i + 1
    rPr.insert(insert_at, effect)


def _apply_text_color_to_text_frame(text_frame, rgb: tuple[int, int, int]) -> None:
    """Apply font color plus a contrast outline/shadow to all runs in a text frame."""
    r, g, b = rgb
    color = RGBColor(r, g, b)
    effect_rgb = _contrast_color(rgb)

    for paragraph in text_frame.paragraphs:
        # If a paragraph has no runs (rare), python-pptx may not expose font directly.
        # Most shapes will have runs; apply color to each run.
        for run in paragraph.runs:
            run.font.color.rgb = color
            try:
                _apply_outline_and_shadow_to_run(run, effect_rgb)
            except Exception:
                # Never let a styling failure break presentation creation.
                pass


def apply_text_color_to_shape(shape, rgb: tuple[int, int, int]) -> None:
    """Recursively apply text color to a shape (text, tables, groups)."""
    # Group shapes: recurse into children
    if hasattr(shape, "shapes"):
        try:
            for subshape in shape.shapes:
                apply_text_color_to_shape(subshape, rgb)
        except Exception:
            pass

    # Regular text frames
    if hasattr(shape, "has_text_frame") and shape.has_text_frame:
        _apply_text_color_to_text_frame(shape.text_frame, rgb)

    # Tables
    if hasattr(shape, "has_table") and shape.has_table:
        table = shape.table
        for row in table.rows:
            for cell in row.cells:
                _apply_text_color_to_text_frame(cell.text_frame, rgb)


def creat_powerpoint_background(presentation_path,saved, info_pic,remove_only = False):
    prs = Presentation(presentation_path)
    
    for slide in prs.slides:
        # Remove any previously-added pictures, then add the new one
        removed = remove_all_pictures(slide)
        if remove_only:
            for shape in slide.shapes:
                apply_text_color_to_shape(shape, (0,0,0))
            continue
        pic = slide.shapes.add_picture(
            CUR_DIR + '/background_images/' +info_pic['image'],
            info_pic['horizontal offset'],
            info_pic['vertical offset'],
            width=info_pic['width'],
            height=info_pic['height'],
        )
        send_shape_to_back(slide, pic)

        # Re-color all text so lyrics stay readable on top of the background
        for shape in slide.shapes:
            apply_text_color_to_shape(shape, info_pic['color'])

    prs.save(saved)
    saved = saved.split('/')
    if __name__ == '__main__':
        print(f"Background: {saved[-1][:-5]}")


if __name__ == '__main__':
    sys.stdout = open(CUR_DIR + "/songs done.txt", "w")   # "w" overwrites the file each run
    i = 0
    for presentation in os.listdir(CUR_DIR + '/To Add Background'):
        if presentation[-5:] == '.pptx':
            prs = Presentation(CUR_DIR + '/To Add Background/' + presentation)
            info_pic = Background_image(CUR_DIR + '/background_images',prs)
            creat_powerpoint_background(CUR_DIR + '/To Add Background/' + presentation,OUTPUT_PPTX + presentation, info_pic,remove_only=False)
            i += 1
    print('Presentation created =',i)
    print()
    sys.stdout.close()
