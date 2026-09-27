"""HTTP integration tests for the web and desktop presentation workflows."""
import io
import json
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from urllib.request import Request, urlopen
import zipfile

from pptx import Presentation
import app


class PresentationWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings = patch.object(app.core, 'SETTINGS_FILE', str(Path(self.temp.name) / 'settings.json'))
        self.settings.start()
        self.addCleanup(self.settings.stop)
        self.server = app.ThreadingHTTPServer(('127.0.0.1', 0), app.Handler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.base = f'http://127.0.0.1:{self.server.server_port}'
        self.addCleanup(self.stop_server)
        self.cloud = hasattr(app, '_DOWNLOADS')
        if self.cloud:
            for name, value in [('_DOWNLOAD_DIR', self.temp.name), ('_DOWNLOADS', {})]:
                p = patch.object(app, name, value)
                p.start()
                self.addCleanup(p.stop)

    def stop_server(self):
        self.server.shutdown()
        self.server.server_close()
        self.thread.join()

    def post(self, path, payload):
        req = Request(self.base + path, data=json.dumps(payload).encode(),
                      headers={'Content-Type': 'application/json'})
        with urlopen(req, timeout=30) as response:
            result = json.load(response)
        self.assertTrue(result.get('ok'), result)
        return result

    def create(self, line='First original lyric line'):
        return self.post('/api/create_one', {
            'title': 'Workflow Test', 'artist': 'Test Artist',
            'lyrics': line + '\nSecond original lyric line',
            'background': 'Book.jpeg', 'output_folder': self.temp.name,
            'storage_mode': 'output_only', 'skip_output': True,
            'previous_powerpoints_action': 'keep',
        })

    def read_deck(self, result):
        if self.cloud:
            with urlopen(self.base + result['download'], timeout=30) as response:
                data = response.read()
        else:
            data = Path(result['path']).read_bytes()
        presentation = Presentation(io.BytesIO(data))
        self.assertGreaterEqual(len(presentation.slides), 2)
        return '\n'.join(shape.text for slide in presentation.slides
                         for shape in slide.shapes if shape.has_text_frame)

    def test_interface_backgrounds_and_powerpoint(self):
        with urlopen(self.base, timeout=30) as response:
            self.assertIn('Lyrics', response.read().decode())
        with urlopen(self.base + '/api/backgrounds', timeout=30) as response:
            backgrounds = json.load(response)
        self.assertTrue(backgrounds['ok'])
        self.assertGreater(len(backgrounds['backgrounds']), 0)
        self.assertIn('First original lyric line', self.read_deck(self.create()))

    def test_same_title_downloads_remain_independent_and_zip_is_valid(self):
        if not self.cloud:
            self.skipTest('Web download workflow')
        first = self.create()
        second = self.create('Replacement original lyric line')
        self.assertNotEqual(first['download'], second['download'])
        self.assertIn('First original lyric line', self.read_deck(first))
        self.assertIn('Replacement original lyric line', self.read_deck(second))
        bundle = self.post('/api/download_all', {
            'tokens': [r['download'].rsplit('/', 1)[1] for r in (first, second)],
            'name': 'Sunday Set',
        })
        with urlopen(self.base + bundle['download'], timeout=30) as response:
            with zipfile.ZipFile(io.BytesIO(response.read())) as archive:
                self.assertEqual(2, len(archive.namelist()))
                self.assertIsNone(archive.testzip())


if __name__ == '__main__':
    unittest.main()
